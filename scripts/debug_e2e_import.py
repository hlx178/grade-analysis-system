from io import BytesIO

from openpyxl import Workbook, load_workbook

from app import create_app


def print_kv(k, v):
    try:
        print(f"{k}: {v}")
    except Exception:
        print(k, v)


def main():
    app = create_app()
    with app.app_context():
        c = app.test_client()
        # login admin
        r = c.post("/auth/login", data={"username": "admin", "password": "123456"})
        print_kv("login", r.status_code)
        # 1) download grades template and show header
        rtpl = c.get("/api/import/template/grades?v=2")
        print_kv("template_status", rtpl.status_code)
        try:
            wb = load_workbook(BytesIO(rtpl.data))
            ws = wb.active
            header = [cell.value for cell in ws[1]]
            print_kv("template_header", header)
        except Exception as e:
            print_kv("template_error", str(e))
        # 2) build a minimal workbook for import
        wb2 = Workbook()
        ws2 = wb2.active
        ws2.title = "Sheet1"
        ws2.append(["学号", "姓名", "班级", "语文", "数学"])
        ws2.append(["", "张三", "高一(1)班", 90, 88])
        ws2.append(["20240002", "李四", "高一(1)班", 85, 81])
        buf = BytesIO()
        wb2.save(buf)
        buf.seek(0)
        # 3) preview upload
        rp = c.post(
            "/api/import/preview",
            data={"file": (buf, "grades_test.xlsx")},
            content_type="multipart/form-data",
        )
        print_kv("preview", rp.status_code)
        pj = rp.get_json() or {}
        print_kv("preview_json", pj)
        fid = pj.get("file_id")
        sheets = pj.get("sheets") or []
        sh = sheets[0] if sheets else None
        # 4) inspect
        ri = c.post("/api/import/preview/inspect", json={"file_id": fid, "sheet_name": sh})
        print_kv("inspect", ri.status_code)
        ij = ri.get_json() or {}
        print_kv("inspect_json", ij)
        # 5) import grades
        rg = c.post(
            "/api/import/grades",
            json={
                "file_id": fid,
                "sheet_name": sh,
                "exam_name": "测试模板v2",
                "exam_type": "regular",
            },
        )
        print_kv("import", rg.status_code)
        try:
            print_kv("import_json", rg.get_json())
        except Exception:
            print_kv("import_text", rg.data.decode("utf-8", "ignore"))
        # 6) aggregated check
        ra = c.get("/api/grades/aggregated?exam_name=" + "测试模板v2")
        print_kv("aggregated", ra.status_code)
        try:
            aj = ra.get_json()
            if isinstance(aj, dict):
                print_kv("aggregated_keys", list(aj.keys()))
            else:
                print_kv("aggregated_type", type(aj))
        except Exception:
            print_kv("aggregated_text", ra.data.decode("utf-8", "ignore"))


if __name__ == "__main__":
    main()
